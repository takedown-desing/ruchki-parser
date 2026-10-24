"""Excel в формате импортёра сайта ручкидлядверей.рф: строка = покрытие, строки одной модели идут подряд с одинаковым «Названием».
Колонки «Источник: …» служебные: импортёр их пропускает, они нужны для проверки."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_COLS = ['Название', 'Бренд', 'Раздел', 'Артикул', 'Покрытие', 'Артикул варианта', 'Цена', 'На складе', 'Фото', 'Краткое описание']
SRC_COLS = ['Источник: адрес', 'Источник: раздел', 'Источник: описание', 'Источник: все фото', 'Источник: фото в репозитории', 'Источник: цена не в рублях']
HEAD_FILL = PatternFill('solid', fgColor='F3EEDD')


def section_for(model, mapping):
    """Раздел нашего сайта по разделу источника. Пусто, если соответствие не задано: импортёр тогда не создаст лишний раздел."""
    cat = model.get('category') or []
    if not cat or not mapping:
        return ''
    full = ' > '.join(cat)
    return mapping.get(full) or mapping.get(cat[-1]) or ''


def rows_for(models, site, images_prefix=''):
    spec_keys = []
    for m in models:
        for k in m['specs']:
            if k not in spec_keys:
                spec_keys.append(k)
    head = BASE_COLS + ['Х: ' + k for k in spec_keys] + SRC_COLS
    rows = []
    for m in models:
        for i, v in enumerate(m['variants']):
            first = i == 0
            rub = (v.get('currency') or 'RUB').upper() in ('RUB', 'RUR', '')
            price = v.get('price')
            stock = '' if v.get('in_stock') is None else ('да' if v['in_stock'] else 'нет')
            row = {
                'Название': m['name'],
                'Бренд': m['brand'] if first else '',
                'Раздел': section_for(m, site.get('mapping') or {}) if first else '',
                'Артикул': m['sku'] if first else '',
                'Покрытие': v['finish'],
                'Артикул варианта': v['sku'],
                'Цена': (int(price) if price and float(price).is_integer() else price) if (price and rub) else '',
                'На складе': stock,
                'Фото': v['images'][0] if v['images'] else '',
                'Краткое описание': (m['description'][:600] if site.get('use_description') else '') if first else '',
                'Источник: адрес': v['url'],
                'Источник: раздел': ' > '.join(m.get('category') or []) if first else '',
                'Источник: описание': m['description'][:3000] if first else '',
                'Источник: все фото': '\n'.join(v['images'][:12]),
                'Источник: фото в репозитории': '\n'.join(images_prefix + x for x in v.get('local_images', [])),
                'Источник: цена не в рублях': f"{price} {v.get('currency')}" if (price and not rub) else '',
            }
            for k in spec_keys:
                row['Х: ' + k] = m['specs'].get(k, '') if first else ''
            rows.append([row.get(h, '') for h in head])
    return head, rows


def write_products(path, models, site, images_prefix=''):
    head, rows = rows_for(models, site, images_prefix)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Товары'
    ws.append(head)
    for r in rows:
        ws.append(r)
    _style(ws, head)
    wb.save(path)
    return len(rows)


def write_changes(path, changes):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Изменения'
    head = ['Что', 'Название', 'Покрытие', 'Артикул варианта', 'Было', 'Стало']
    ws.append(head)
    labels = {'new': 'новое', 'gone': 'пропало из источника', 'price': 'цена', 'stock': 'наличие'}
    fmt = lambda x: '' if x is None else ('да' if x is True else 'нет' if x is False else x)
    for kind in ('new', 'price', 'stock', 'gone'):
        for c in changes.get(kind, []):
            was = c.get('old') if kind in ('price', 'stock') else (c.get('price') if kind == 'gone' else '')
            now = c.get('new') if kind in ('price', 'stock') else (c.get('price') if kind == 'new' else '')
            ws.append([labels[kind], c['name'], c['finish'], c['sku'], fmt(was), fmt(now)])
    _style(ws, head)
    wb.save(path)


def _style(ws, head):
    for i, h in enumerate(head, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical='top', wrap_text=True)
        width = 14 if h.startswith('Х: ') else 22
        if h in ('Название', 'Источник: адрес', 'Источник: описание', 'Фото', 'Источник: все фото'):
            width = 40
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'B2'
