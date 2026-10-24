"""Проверки без интернета: разбор фидов, извлечение товара со страницы, склейка покрытий, Excel."""
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from parser import feeds, pages, excel  # noqa: E402
from parser.models import build_models, diff  # noqa: E402
from parser.item import split_color, to_price, new_item  # noqa: E402

YML = '''<?xml version="1.0" encoding="UTF-8"?>
<yml_catalog date="2026-10-24"><shop><categories>
<category id="1">Дверные ручки</category><category id="2" parentId="1">На розетке</category></categories>
<offers>
<offer id="10" available="true" group_id="77"><name>Ручка Delta матовый чёрный</name><vendor>TUPAI</vendor><vendorCode>D-01-BL</vendorCode>
<url>https://x.ru/p/10</url><price>4990</price><currencyId>RUR</currencyId><categoryId>2</categoryId>
<picture>https://x.ru/i/10.jpg</picture><picture>https://x.ru</picture><description>&lt;p&gt;Латунная ручка&lt;/p&gt;</description>
<param name="Материал">латунь</param><param name="Цвет">матовый чёрный</param></offer>
<offer id="11" group_id="77"><name>Ручка Delta матовый хром</name><vendor>TUPAI</vendor><vendorCode>D-01-CR</vendorCode>
<price>5190</price><categoryId>2</categoryId><availability>ACTIVE</availability><param name="Цвет">матовый хром</param></offer>
</offers></shop></yml_catalog>'''

GOOGLE = '''<?xml version="1.0"?><rss xmlns:g="http://base.google.com/ns/1.0" version="2.0"><channel>
<item><g:id>A1</g:id><g:title>Knob Alfa</g:title><g:price>1200.00 RUB</g:price><g:availability>in stock</g:availability>
<g:image_link>https://y.ru/a.jpg</g:image_link><g:brand>Fimet</g:brand><g:product_type>Ручки &gt; Кнобы</g:product_type></item>
</channel></rss>'''

PAGE = '''<html><head><script type="application/ld+json">{"@context":"https://schema.org","@graph":[
{"@type":"BreadcrumbList","itemListElement":[{"position":1,"name":"Главная"},{"position":2,"name":"Ручки"},{"position":3,"name":"На розетке"}]},
{"@type":"Product","name":"Ручка Omega цвет бронза","sku":"OM-1.BR","brand":{"@type":"Brand","name":"Colombo"},
 "image":["/img/om1.jpg","https://z.ru/img/om2.jpg"],"description":"Ручка на розетке.",
 "offers":{"@type":"Offer","price":"7900","priceCurrency":"RUB","availability":"https://schema.org/InStock"}}]}</script></head>
<body><h1>Ручка Omega цвет бронза</h1><table><tr><td>Материал</td><td>латунь</td></tr><tr><td>Страна</td><td>Италия</td></tr></table></body></html>'''


def test_detect_and_yml():
    assert feeds.detect(YML) == 'yml'
    items = feeds.parse(YML, 'yml')
    assert len(items) == 2
    a, b = items
    assert a['category'] == ['Дверные ручки', 'На розетке']
    assert a['price'] == 4990 and a['in_stock'] is True and a['finish'] == 'матовый чёрный'
    assert a['images'] == ['https://x.ru/i/10.jpg'], 'пустая картинка-главная отброшена'
    assert a['specs'] == {'Материал': 'латунь'} and a['description'] == 'Латунная ручка'
    assert b['in_stock'] is True, 'availability ACTIVE = в наличии'


def test_google_feed():
    assert feeds.detect(GOOGLE) == 'google'
    it = feeds.parse(GOOGLE, 'google')[0]
    assert it['name'] == 'Knob Alfa' and it['price'] == 1200 and it['currency'] == 'RUB'
    assert it['in_stock'] is True and it['category'] == ['Ручки', 'Кнобы'] and it['brand'] == 'Fimet'


def test_not_a_feed():
    assert feeds.detect('<html><body>Каталог</body></html>') is None
    assert feeds.detect('') is None


def test_page_jsonld():
    assert pages.is_product_page(PAGE)
    it = pages.extract(PAGE, 'https://z.ru/p/omega')
    assert it['name'] == 'Ручка Omega цвет бронза' and it['sku'] == 'OM-1.BR' and it['brand'] == 'Colombo'
    assert it['price'] == 7900 and it['in_stock'] is True
    assert it['images'][0] == 'https://z.ru/img/om1.jpg', 'относительный адрес фото стал полным'
    assert it['category'] == ['Ручки', 'На розетке']
    assert it['specs']['Материал'] == 'латунь' and it['specs']['Страна'] == 'Италия'


def test_split_color_and_price():
    assert split_color('Ручка X на розетке цвет бронза') == ('Ручка X на розетке', 'бронза')
    assert split_color('Ручка X матовый чёрный') == ('Ручка X', 'матовый чёрный')
    assert split_color('Упор дверной 45') == ('Упор дверной 45', '')
    assert to_price('от 5 490 руб.') == 5490 and to_price('нет') is None and to_price('1200.50 RUB') == 1200.5


def _it(name, sku, finish='', group=''):
    i = new_item('t'); i.update(name=name, sku=sku, finish=finish, brand='B', group=group); return i


def test_merge_models():
    ms = build_models([_it('Ручка Colombo ONE CC11RSB.C01', 'CC11RSB.C01', 'белый'), _it('Ручка Colombo ONE CC11RSB.C02', 'CC11RSB.C02', 'бронза'),
                       _it('Ручка Tupai 4213 5S 158', '4213 5S 158'), _it('Ручка Tupai 4213 5S 03', '4213 5S 03'),
                       _it('Упор дверной 45', '45')])
    names = {m['name']: m for m in ms}
    assert names['Ручка Colombo ONE CC11RSB']['sku'] == 'CC11RSB'
    assert [v['finish'] for v in names['Ручка Colombo ONE CC11RSB']['variants']] == ['белый', 'бронза']
    assert [v['finish'] for v in names['Ручка Tupai 4213 5S']['variants']] == ['158', '03']
    assert 'Упор дверной 45' in names, 'одиночный товар с числом в конце не обрезан'


def test_merge_by_feed_group():
    ms = build_models(feeds.parse(YML, 'yml'))
    assert len(ms) == 1 and ms[0]['name'] == 'Ручка Delta' and len(ms[0]['variants']) == 2


def test_diff():
    old = build_models(feeds.parse(YML, 'yml'))
    new = build_models(feeds.parse(YML.replace('<price>4990', '<price>5500'), 'yml'))
    d = diff(old, new)
    assert d['price'] and d['price'][0]['old'] == 4990 and d['price'][0]['new'] == 5500 and not d['new'] and not d['gone']


def test_excel_columns(tmp_path):
    ms = build_models(feeds.parse(YML, 'yml'))
    p = tmp_path / 'x.xlsx'
    excel.write_products(str(p), ms, {'mapping': {'Дверные ручки > На розетке': '/catalog/dvernye-ruchki/na-rozetke/'}})
    ws = load_workbook(p).active
    head = [c.value for c in ws[1]]
    row = dict(zip(head, [c.value for c in ws[2]]))
    row2 = dict(zip(head, [c.value for c in ws[3]]))
    assert head[:3] == ['Название', 'Бренд', 'Раздел']
    assert row['Раздел'] == '/catalog/dvernye-ruchki/na-rozetke/' and row['Цена'] == 4990 and row['На складе'] == 'да'
    assert row2['Название'] == row['Название'] and row2['Бренд'] in (None, ''), 'строки одной модели: одно название, поля модели в первой'
    assert row['Х: Материал'] == 'латунь' and row['Краткое описание'] in (None, ''), 'описание в карточку только по галочке'


def test_merge_code_before_color():
    ms = build_models([_it('Ручка Colombo ONE CC11RSB.C01 белый матовый', 'CC11RSB.C01', 'белый матовый'),
                       _it('Ручка Colombo ONE CC11RSB.C02 цвет бронза', 'CC11RSB.C02', 'бронза матовая')])
    assert len(ms) == 1 and ms[0]['name'] == 'Ручка Colombo ONE CC11RSB'
    assert [v['finish'] for v in ms[0]['variants']] == ['белый матовый', 'бронза матовая']


def test_diagnose_reasons():
    from parser import diagnose
    st = {'candidates': 0}
    assert 'не найдено ссылок' in diagnose.empty_reason(st)
    assert 'Поиск на сайте' in diagnose.empty_reason(st, 'AGB', is_search=True)
    it = new_item('page'); it['name'] = 'Ама'
    st = {'candidates': 5, 'samples': [('u', '<html><iframe src="https://x/3D/visore.html"></iframe>' + 'текст ' * 200 + '</html>', it)]}
    assert '3D-просмотрщике' in diagnose.empty_reason(st)
    st = {'candidates': 5, 'samples': [('u', '<html>' + 'текст ' * 200 + '</html>', it)]}
    assert 'фото на них нет' in diagnose.empty_reason(st)


def test_diagnose_block_page():
    from parser import diagnose

    class FakeFx:
        last_status, last_body, last_exc, last_headers = 429, '<html>reCaptcha</html>', None, 'x-ips-suspection: BlockedRegion'

        def get(self, *a, **k):
            return None
    assert 'региональная блокировка' in diagnose.access_problem(FakeFx(), 'https://x.by/')
